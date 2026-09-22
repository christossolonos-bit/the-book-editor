"""Export edited manuscript to a 6x9 hardcopy/print Word file."""

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parent
JSON_PATH = ROOT / "manuscript-export.json"
OUT_PATH = ROOT / "Prove Them Wrong - 6x9 Hardcopy.docx"


def ensure_style(styles, name, style_type=WD_STYLE_TYPE.PARAGRAPH):
    try:
        return styles.add_style(name, style_type)
    except Exception:
        return styles[name]


def configure_styles(doc):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.2
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.first_line_indent = Inches(0.25)

    title_style = ensure_style(styles, "BookTitle")
    title_style.font.name = "Times New Roman"
    title_style.font.size = Pt(26)
    title_style.font.bold = True
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.space_after = Pt(12)
    title_style.paragraph_format.first_line_indent = Inches(0)

    sub_style = ensure_style(styles, "BookSubtitle")
    sub_style.font.name = "Times New Roman"
    sub_style.font.size = Pt(13)
    sub_style.font.italic = True
    sub_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_style.paragraph_format.space_after = Pt(18)
    sub_style.paragraph_format.first_line_indent = Inches(0)

    author_style = ensure_style(styles, "BookAuthor")
    author_style.font.name = "Times New Roman"
    author_style.font.size = Pt(14)
    author_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author_style.paragraph_format.first_line_indent = Inches(0)

    chap_style = ensure_style(styles, "ChapterTitle")
    chap_style.font.name = "Times New Roman"
    chap_style.font.size = Pt(16)
    chap_style.font.bold = True
    chap_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    chap_style.paragraph_format.space_before = Pt(0)
    chap_style.paragraph_format.space_after = Pt(18)
    chap_style.paragraph_format.first_line_indent = Inches(0)

    part_style = ensure_style(styles, "PartLabel")
    part_style.font.name = "Times New Roman"
    part_style.font.size = Pt(11)
    part_style.font.bold = True
    part_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    part_style.paragraph_format.space_after = Pt(8)
    part_style.paragraph_format.first_line_indent = Inches(0)

    body_style = ensure_style(styles, "BodyTextBook")
    body_style.font.name = "Times New Roman"
    body_style.font.size = Pt(12)
    body_style.paragraph_format.line_spacing = 1.2
    body_style.paragraph_format.space_after = Pt(0)
    body_style.paragraph_format.first_line_indent = Inches(0.25)

    first_style = ensure_style(styles, "BodyFirst")
    first_style.font.name = "Times New Roman"
    first_style.font.size = Pt(12)
    first_style.paragraph_format.line_spacing = 1.2
    first_style.paragraph_format.space_after = Pt(0)
    first_style.paragraph_format.first_line_indent = Inches(0)

    section_style = ensure_style(styles, "SectionHead")
    section_style.font.name = "Times New Roman"
    section_style.font.size = Pt(12)
    section_style.font.bold = True
    section_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    section_style.paragraph_format.space_before = Pt(6)
    section_style.paragraph_format.space_after = Pt(10)
    section_style.paragraph_format.first_line_indent = Inches(0)


def setup_page(doc):
    for section in doc.sections:
        section.page_width = Inches(6)
        section.page_height = Inches(9)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.65)
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.gutter = Inches(0.15)


def add_page_break(paragraph):
    """Mark paragraph to start on a new page."""
    pPr = paragraph._p.get_or_add_pPr()
    pageBreak = OxmlElement("w:pageBreakBefore")
    pPr.append(pageBreak)


def add_answer_lines(doc, count=14):
    label = doc.add_paragraph("Your answer", style="SectionHead")
    label.paragraph_format.space_before = Pt(12)
    for _ in range(count):
        p = doc.add_paragraph("", style="BodyFirst")
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "12")
        bottom.set(qn("w:color"), "AAAAAA")
        pBdr.append(bottom)
        pPr.append(pBdr)
        p.paragraph_format.space_after = Pt(10)


def add_chapter_body(doc, text):
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    blocks = [b.strip() for b in text.split("\n\n") if b.strip()]
    first_body_done = False

    for block in blocks:
        lines = [ln.strip() for ln in block.split("\n") if ln.strip()]
        if not lines:
            continue

        # Section starters get their own page in hardcopy
        if re.match(r"^Question\s+\d+$", lines[0], re.I) or re.match(
            r"^(Action step|Final words|Quick last note)\.?$", lines[0], re.I
        ) or re.match(r"^Part\s+\d+\s*:", lines[0], re.I):
            head = doc.add_paragraph(lines[0], style="SectionHead")
            add_page_break(head)
            rest = lines[1:]
            if rest:
                short = all(len(ln) < 60 for ln in rest) and len(rest) > 1
                if short:
                    for ln in rest:
                        doc.add_paragraph(ln, style="BodyFirst")
                else:
                    doc.add_paragraph(" ".join(rest), style="BodyFirst")
            if re.match(r"^Question\s+\d+$", lines[0], re.I):
                add_answer_lines(doc, 14)
            continue

        short_lines = all(len(ln) < 60 for ln in lines) and len(lines) > 1
        if short_lines:
            for ln in lines:
                doc.add_paragraph(ln, style="BodyFirst")
        else:
            para_text = " ".join(lines)
            style = "BodyFirst" if not first_body_done else "BodyTextBook"
            doc.add_paragraph(para_text, style=style)
            first_body_done = True


def main():
    with open(JSON_PATH, encoding="utf8") as f:
        book = json.load(f)

    doc = Document()
    setup_page(doc)
    configure_styles(doc)

    for _ in range(8):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0)

    doc.add_paragraph(book["title"], style="BookTitle")
    doc.add_paragraph(book["subtitle"], style="BookSubtitle")
    doc.add_paragraph(book.get("author", "Christos Solonos"), style="BookAuthor")

    last_part = None
    for ch in book["chapters"]:
        doc.add_page_break()

        part = ch.get("part")
        if part and part != last_part and part not in ("Front", "Close"):
            doc.add_paragraph(part.upper(), style="PartLabel")
        last_part = part

        if ch.get("number") is not None:
            doc.add_paragraph(f"Chapter {ch['number']}", style="ChapterTitle")
            doc.add_paragraph(ch["title"], style="ChapterTitle")
        else:
            doc.add_paragraph(ch["title"], style="ChapterTitle")

        add_chapter_body(doc, ch.get("text"))

    doc.save(OUT_PATH)
    print(f"Saved: {OUT_PATH}")
    print("Trim: 6 in x 9 in")
    print(f"Chapters: {len(book['chapters'])}")


if __name__ == "__main__":
    main()
