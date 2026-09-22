"""Export edited manuscript to a 5x8 Kindle/KDP Word file."""

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parent
JSON_PATH = ROOT / "manuscript-export.json"
OUT_PATH = ROOT / "Prove Them Wrong - 5x8 Kindle.docx"


def ensure_style(styles, name, style_type=WD_STYLE_TYPE.PARAGRAPH):
    try:
        return styles.add_style(name, style_type)
    except Exception:
        return styles[name]


def configure_styles(doc):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.first_line_indent = Inches(0.2)

    title_style = ensure_style(styles, "BookTitle")
    title_style.font.name = "Times New Roman"
    title_style.font.size = Pt(22)
    title_style.font.bold = True
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.space_after = Pt(12)
    title_style.paragraph_format.first_line_indent = Inches(0)

    sub_style = ensure_style(styles, "BookSubtitle")
    sub_style.font.name = "Times New Roman"
    sub_style.font.size = Pt(12)
    sub_style.font.italic = True
    sub_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_style.paragraph_format.space_after = Pt(18)
    sub_style.paragraph_format.first_line_indent = Inches(0)

    author_style = ensure_style(styles, "BookAuthor")
    author_style.font.name = "Times New Roman"
    author_style.font.size = Pt(14)
    author_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author_style.paragraph_format.first_line_indent = Inches(0)

    chap_style = ensure_style(styles, "ChapterTitle")
    chap_style.font.name = "Times New Roman"
    chap_style.font.size = Pt(14)
    chap_style.font.bold = True
    chap_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    chap_style.paragraph_format.space_before = Pt(0)
    chap_style.paragraph_format.space_after = Pt(18)
    chap_style.paragraph_format.first_line_indent = Inches(0)

    part_style = ensure_style(styles, "PartLabel")
    part_style.font.name = "Times New Roman"
    part_style.font.size = Pt(10)
    part_style.font.bold = True
    part_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    part_style.paragraph_format.space_after = Pt(6)
    part_style.paragraph_format.first_line_indent = Inches(0)

    body_style = ensure_style(styles, "BodyTextBook")
    body_style.font.name = "Times New Roman"
    body_style.font.size = Pt(11)
    body_style.paragraph_format.line_spacing = 1.15
    body_style.paragraph_format.space_after = Pt(0)
    body_style.paragraph_format.first_line_indent = Inches(0.2)

    first_style = ensure_style(styles, "BodyFirst")
    first_style.font.name = "Times New Roman"
    first_style.font.size = Pt(11)
    first_style.paragraph_format.line_spacing = 1.15
    first_style.paragraph_format.space_after = Pt(0)
    first_style.paragraph_format.first_line_indent = Inches(0)


def setup_page(doc):
    for section in doc.sections:
        section.page_width = Inches(5)
        section.page_height = Inches(8)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.5)
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.gutter = Inches(0.1)


def add_chapter_body(doc, text):
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    blocks = [b.strip() for b in text.split("\n\n") if b.strip()]

    for i, block in enumerate(blocks):
        lines = [ln.strip() for ln in block.split("\n") if ln.strip()]
        short_lines = all(len(ln) < 60 for ln in lines) and len(lines) > 1

        if short_lines:
            for j, ln in enumerate(lines):
                p = doc.add_paragraph(ln, style="BodyFirst")
                p.paragraph_format.space_after = Pt(2)
                if i == 0 and j == 0:
                    p.paragraph_format.space_before = Pt(0)
        else:
            para_text = " ".join(lines)
            style = "BodyFirst" if i == 0 else "BodyTextBook"
            doc.add_paragraph(para_text, style=style)


def main():
    with open(JSON_PATH, encoding="utf8") as f:
        book = json.load(f)

    doc = Document()
    setup_page(doc)
    configure_styles(doc)

    # Title page
    for _ in range(6):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0)

    doc.add_paragraph(book["title"], style="BookTitle")
    doc.add_paragraph(book["subtitle"], style="BookSubtitle")
    doc.add_paragraph(book.get("author", ""), style="BookAuthor")

    last_part = None
    for ch in book["chapters"]:
        doc.add_page_break()

        part = ch.get("part")
        if part and part != last_part and part not in ("Front", "Close"):
            doc.add_paragraph(part.upper(), style="PartLabel")
        last_part = part

        if ch.get("number") is not None:
            doc.add_paragraph(f"Chapter {ch['number']}", style="ChapterTitle")
            doc.add_paragraph(ch["title"], style="ChapterTitle")
        else:
            doc.add_paragraph(ch["title"], style="ChapterTitle")

        add_chapter_body(doc, ch.get("text"))

    doc.save(OUT_PATH)
    print(f"Saved: {OUT_PATH}")
    print("Trim: 5 in x 8 in")
    print(f"Chapters: {len(book['chapters'])}")


if __name__ == "__main__":
    main()
