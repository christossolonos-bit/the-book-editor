"""
Full book export for Google Docs upload.
Includes cover, chapter images, formatted text, part breaks, and answer lines.

Upload the .docx to Google Drive → Open with Google Docs.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from export_layout import chapter_pages, is_shout as layout_is_shout

ROOT = Path(__file__).resolve().parent
JSON_PATH = ROOT / "manuscript-export.json"
ART_DIR = ROOT / "public" / "art"
COVER = ROOT / "public" / "cover.png"
OUT_DIR = ROOT / "google-docs-export"
AMAZON_DIR = ROOT / "amazon-ready"

FORMATS = {
    "print": {
        "filename": "Prove Them Wrong - 6x9 Hardcopy (Google Docs).docx",
        "width": 6,
        "height": 9,
        "margins": (0.75, 0.65, 0.75, 0.75),
        "body_pt": 12,
        "title_pt": 26,
        "chap_pt": 16,
        "line_spacing": 1.2,
        "indent": 0.25,
        "answer_lines": 14,
        "image_width": 3.6,
        "image_max_height": 2.4,
        "cover_width": 5.4,
    },
    "kindle": {
        "filename": "Prove Them Wrong - 5x8 Kindle (Google Docs).docx",
        "width": 5,
        "height": 8,
        "margins": (0.6, 0.5, 0.6, 0.6),
        "body_pt": 11,
        "title_pt": 22,
        "chap_pt": 14,
        "line_spacing": 1.15,
        "indent": 0.2,
        "answer_lines": 12,
        "image_width": 3.0,
        "image_max_height": 2.1,
        "cover_width": 4.4,
    },
}


def ensure_style(styles, name):
    try:
        return styles.add_style(name, 1)
    except Exception:
        return styles[name]


def set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def page_break_before(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    el = OxmlElement("w:pageBreakBefore")
    pPr.append(el)


def add_bottom_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "10")
    bottom.set(qn("w:color"), "AAAAAA")
    pBdr.append(bottom)
    pPr.append(pBdr)


def configure_doc(doc, fmt):
    for section in doc.sections:
        section.page_width = Inches(fmt["width"])
        section.page_height = Inches(fmt["height"])
        left, right, top, bottom = fmt["margins"]
        section.left_margin = Inches(left)
        section.right_margin = Inches(right)
        section.top_margin = Inches(top)
        section.bottom_margin = Inches(bottom)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(fmt["body_pt"])
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = fmt["line_spacing"]
    pf.space_after = Pt(0)
    pf.first_line_indent = Inches(fmt["indent"])


def add_centered_para(doc, text, size, bold=False, italic=False, space_after=8, space_before=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_body_para(doc, text, fmt, first=False, centered=False, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = fmt["line_spacing"]
    p.paragraph_format.space_after = Pt(0)
    if centered:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
    elif first:
        p.paragraph_format.first_line_indent = Inches(0)
    else:
        p.paragraph_format.first_line_indent = Inches(fmt["indent"])
    run = p.add_run(text)
    set_run_font(run, size=fmt["body_pt"], bold=bold)
    return p


def add_answer_lines(doc, count):
    p = add_centered_para(doc, "Your answer", size=11, bold=True, space_before=14, space_after=8)
    for _ in range(count):
        line = doc.add_paragraph()
        line.paragraph_format.first_line_indent = Inches(0)
        line.paragraph_format.space_after = Pt(8)
        add_bottom_border(line)


def add_image_centered(doc, path: Path, width_in: float, max_height_in: float | None = None):
    if not path or not Path(path).exists():
        return
    path = Path(path)
    width = float(width_in)
    if max_height_in:
        from PIL import Image as PILImage

        with PILImage.open(path) as im:
            w, h = im.size
        aspect = h / float(w) if w else 1.0
        if width * aspect > max_height_in:
            width = max_height_in / aspect
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))


def emit_block(doc, block: dict, fmt: dict, first_body: list):
    text = block["text"]
    btype = block["type"]

    if btype == "section":
        add_centered_para(doc, text, size=fmt["body_pt"], bold=True, space_before=6, space_after=10)
        return

    if re.match(r"^Try this:", text, re.I):
        add_body_para(doc, text, fmt, first=True, centered=True, bold=True)
        return

    if btype == "shout" or layout_is_shout(text):
        add_body_para(doc, text, fmt, first=True, centered=True, bold=True)
        return

    add_body_para(doc, text, fmt, first=not first_body[0])
    first_body[0] = True


def build(fmt_key: str, book: dict) -> Path:
    fmt = FORMATS[fmt_key]
    doc = Document()
    configure_doc(doc, fmt)

    if COVER.exists():
        add_image_centered(doc, COVER, fmt["cover_width"])
    else:
        for _ in range(6):
            doc.add_paragraph()
        add_centered_para(doc, book["title"], fmt["title_pt"], bold=True, space_after=12)
        add_centered_para(doc, book["subtitle"], fmt["body_pt"], italic=True, space_after=18)
        add_centered_para(doc, book.get("author", "Christos Solonos"), fmt["body_pt"] + 1)

    last_part = None
    for ch in book["chapters"]:
        pages = chapter_pages(ch, fmt_key)
        first_body = [False]

        for page in pages:
            # Each export page starts on a new Word page
            if page["is_first"]:
                doc.add_page_break()
                part = ch.get("part")
                if part and part != last_part and part not in ("Front", "Close"):
                    add_centered_para(doc, part.upper(), 10, bold=True, space_after=8)
                last_part = part

                if ch.get("number") is not None:
                    add_centered_para(doc, f"Chapter {ch['number']}", 11, bold=False, space_after=4)
                    add_centered_para(doc, ch["title"], fmt["chap_pt"], bold=True, space_after=16)
                else:
                    add_centered_para(doc, ch["title"], fmt["chap_pt"], bold=True, space_after=16)
                add_centered_para(doc, "❦", 12, space_after=14)
            else:
                # Force page break via empty paragraph with pageBreakBefore
                p = doc.add_paragraph()
                page_break_before(p)
                first_body[0] = False

            is_question_page = any(
                b["type"] == "section" and re.match(r"^Question\s+\d+", b["text"], re.I)
                for b in page["blocks"]
            )

            for block in page["blocks"]:
                emit_block(doc, block, fmt, first_body)

            if is_question_page:
                add_answer_lines(doc, fmt["answer_lines"])

            if page["show_art"] and page["art_path"]:
                # Cap height from spare lines so Word does not kick art to a lonely page
                spare = page.get("spare_lines", 12)
                line_in = (fmt["body_pt"] * fmt["line_spacing"]) / 72.0
                max_h = min(fmt["image_max_height"], max(spare * line_in * 0.85, 1.4))
                add_image_centered(doc, page["art_path"], fmt["image_width"], max_height_in=max_h)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / fmt["filename"]
    doc.save(out)
    return out


def main():
    book = json.loads(JSON_PATH.read_text(encoding="utf8"))
    AMAZON_DIR.mkdir(parents=True, exist_ok=True)
    amazon_names = {
        "print": "Prove Them Wrong - 6x9 Hardcopy.docx",
        "kindle": "Prove Them Wrong - 5x8 Kindle.docx",
    }
    for key in ("print", "kindle"):
        path = build(key, book)
        amazon = AMAZON_DIR / amazon_names[key]
        amazon.write_bytes(path.read_bytes())
        print(f"Saved: {path}")
        print(f"Copied: {amazon}")
    print("Upload these .docx files to Google Drive, then Open with Google Docs.")


if __name__ == "__main__":
    main()
